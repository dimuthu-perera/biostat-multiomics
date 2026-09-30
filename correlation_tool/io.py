from __future__ import annotations

import csv
import hashlib
import math
import posixpath
import re
import zipfile
import xml.etree.ElementTree as ET
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Iterable, Literal

import numpy as np
import pandas as pd
from openpyxl import load_workbook

Orientation = Literal["samples_rows", "samples_columns"]

# Missing tokens apply ONLY to measurement cells. Identifier cells are literal text.
DEFAULT_MISSING_TOKENS = frozenset({"", "NA", "N/A", "NaN", "nan"})
_INT_RE = re.compile(r"^[+-]?\d+$")
_CELL_REF_RE = re.compile(r"^\$?([A-Za-z]+)\$?([1-9]\d*)$")

# Conservative import limits for the scalar reference implementation.
DEFAULT_MAX_FILE_BYTES = 100 * 1024 * 1024
DEFAULT_MAX_RAW_ROWS = 100_000
DEFAULT_MAX_RAW_COLUMNS = 50_000
DEFAULT_MAX_MATRIX_CELLS = 5_000_000
DEFAULT_MAX_XLSX_UNCOMPRESSED_BYTES = 256 * 1024 * 1024
DEFAULT_MAX_XLSX_ENTRY_BYTES = 128 * 1024 * 1024


class MatrixFormatError(ValueError):
    """Raised when an input matrix violates the explicit import contract."""


def sha256_file(path: str | Path) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def _validate_identifier(value: object, *, kind: str, position: str) -> str:
    if not isinstance(value, str):
        raise MatrixFormatError(
            f"{kind} identifier at {position} must be stored as text; got {type(value).__name__}. "
            "This prevents silent identifier conversion (for example 001 -> 1)."
        )
    if value == "":
        raise MatrixFormatError(f"{kind} identifier at {position} is empty.")
    return value


def _reject_duplicates(values: list[str], *, kind: str) -> None:
    seen: set[str] = set()
    dup: list[str] = []
    for value in values:
        if value in seen and value not in dup:
            dup.append(value)
        seen.add(value)
    if dup:
        shown = ", ".join(repr(x) for x in dup[:8])
        more = " ..." if len(dup) > 8 else ""
        raise MatrixFormatError(f"Duplicate {kind} identifiers detected before parsing: {shown}{more}")


def _reject_nonfinite_numeric(fv: float, *, row: int, column: int, original: object) -> None:
    if math.isnan(fv):
        raise MatrixFormatError(
            f"Measurement at row {row}, column {column} parsed as NaN from {original!r}, but it was not an explicitly configured missing token."
        )
    if math.isinf(fv):
        raise MatrixFormatError(f"Infinite measurement at row {row}, column {column} is unsupported: {original!r}.")


def _parse_measurement(value: object, *, row: int, column: int, missing_tokens: frozenset[str]) -> int | float | None:
    if value is None:
        return None
    if isinstance(value, (bool, np.bool_)):
        raise MatrixFormatError(f"Boolean measurement at row {row}, column {column} is unsupported.")
    if isinstance(value, (complex, np.complexfloating)):
        raise MatrixFormatError(f"Complex measurement at row {row}, column {column} is unsupported.")
    if isinstance(value, (int, np.integer)):
        iv = int(value)
        if iv < np.iinfo(np.int64).min or iv > np.iinfo(np.int64).max:
            raise MatrixFormatError(f"Integer measurement at row {row}, column {column} exceeds int64 range.")
        return iv
    if isinstance(value, (float, np.floating)):
        fv = float(value)
        _reject_nonfinite_numeric(fv, row=row, column=column, original=value)
        return fv
    if isinstance(value, str):
        if value in missing_tokens:
            return None
        if _INT_RE.fullmatch(value):
            try:
                iv = int(value)
            except ValueError as exc:
                raise MatrixFormatError(f"Invalid integer measurement at row {row}, column {column}: {value!r}") from exc
            if iv < np.iinfo(np.int64).min or iv > np.iinfo(np.int64).max:
                raise MatrixFormatError(f"Integer measurement at row {row}, column {column} exceeds int64 range.")
            return iv
        try:
            fv = float(value)
        except ValueError as exc:
            raise MatrixFormatError(
                f"Nonnumeric measurement at row {row}, column {column}: {value!r}. "
                "Only explicit numeric values or configured missing tokens are accepted."
            ) from exc
        _reject_nonfinite_numeric(fv, row=row, column=column, original=value)
        # Detect total loss of a nonzero decimal to binary64 zero. Ordinary rounding is
        # accepted; only complete underflow to zero is rejected as an integrity failure.
        if fv == 0.0:
            try:
                dec = Decimal(value)
            except InvalidOperation:
                dec = Decimal(0)
            if dec.is_finite() and dec != 0:
                raise MatrixFormatError(
                    f"Measurement at row {row}, column {column} underflows to binary64 zero: {value!r}. "
                    "Rescale upstream or provide values on a numerically representable scale."
                )
        return fv
    raise MatrixFormatError(
        f"Unsupported measurement type at row {row}, column {column}: {type(value).__name__}."
    )


def _column_series(values: list[int | float | None], name: str) -> pd.Series:
    nonmissing = [v for v in values if v is not None]
    if not nonmissing:
        return pd.Series(pd.array([pd.NA] * len(values), dtype="Float64"), name=name)

    all_int = all(isinstance(v, int) and not isinstance(v, bool) for v in nonmissing)
    if all_int:
        return pd.Series(pd.array(values, dtype="Int64"), name=name)

    # Any mixed floating column requires float64 representation. Refuse exact integer
    # values that cannot be represented by float64; pre-existing float values are kept
    # as the values actually supplied by the researcher.
    for v in nonmissing:
        if isinstance(v, int) and abs(v) > 2**53:
            raise MatrixFormatError(
                f"Column {name!r} mixes floating-point values with integer values above 2**53; "
                "float64 conversion would lose exact integer identity."
            )
    return pd.Series(np.array([np.nan if v is None else float(v) for v in values], dtype=np.float64), name=name)


def _matrix_from_raw_rows(
    rows: list[list[object]],
    *,
    orientation: Orientation,
    missing_tokens: frozenset[str],
) -> pd.DataFrame:
    if not rows:
        raise MatrixFormatError("The file is empty.")
    if len(rows) < 2:
        raise MatrixFormatError("The file must contain a header row and at least one data row.")
    width = len(rows[0])
    if width < 2:
        raise MatrixFormatError("The matrix must contain an ID column and at least one measurement column.")
    for i, row in enumerate(rows, start=1):
        if len(row) != width:
            raise MatrixFormatError(
                f"Ragged matrix: row {i} contains {len(row)} fields; expected {width}."
            )

    column_ids = [_validate_identifier(v, kind="column", position=f"header column {j}") for j, v in enumerate(rows[0][1:], start=2)]
    row_ids = [_validate_identifier(row[0], kind="row", position=f"row {i}") for i, row in enumerate(rows[1:], start=2)]
    _reject_duplicates(column_ids, kind="column")
    _reject_duplicates(row_ids, kind="row")

    parsed_rows: list[list[int | float | None]] = []
    for i, row in enumerate(rows[1:], start=2):
        parsed_rows.append(
            [
                _parse_measurement(v, row=i, column=j, missing_tokens=missing_tokens)
                for j, v in enumerate(row[1:], start=2)
            ]
        )

    columns: dict[str, pd.Series] = {}
    for j, col_id in enumerate(column_ids):
        columns[col_id] = _column_series([r[j] for r in parsed_rows], col_id)
    df = pd.DataFrame(columns)
    df.index = pd.Index(row_ids, dtype="object")

    if orientation == "samples_rows":
        return df
    if orientation == "samples_columns":
        out = df.transpose().copy()
        out.index = pd.Index(column_ids, dtype="object")
        out.columns = pd.Index(row_ids, dtype="object")
        rebuilt: dict[str, pd.Series] = {}
        for feature in out.columns:
            vals = out[feature].tolist()
            normalized = [None if pd.isna(v) else (int(v) if isinstance(v, (int, np.integer)) else float(v)) for v in vals]
            rebuilt[str(feature)] = _column_series(normalized, str(feature))
        return pd.DataFrame({k: v.array for k, v in rebuilt.items()}, index=pd.Index(column_ids, dtype="object"))
    raise ValueError(f"Unsupported orientation: {orientation}")


def _enforce_raw_dimensions(rows: int, columns: int) -> None:
    if rows > DEFAULT_MAX_RAW_ROWS:
        raise MatrixFormatError(f"Input has {rows:,} raw rows, exceeding the scalar-engine import cap of {DEFAULT_MAX_RAW_ROWS:,}.")
    if columns > DEFAULT_MAX_RAW_COLUMNS:
        raise MatrixFormatError(f"Input has {columns:,} raw columns, exceeding the scalar-engine import cap of {DEFAULT_MAX_RAW_COLUMNS:,}.")
    if rows * columns > DEFAULT_MAX_MATRIX_CELLS:
        raise MatrixFormatError(
            f"Input has {rows * columns:,} raw cells, exceeding the scalar-engine import cap of {DEFAULT_MAX_MATRIX_CELLS:,}."
        )


def _read_delimited(path: Path, delimiter: str) -> list[list[object]]:
    rows: list[list[object]] = []
    try:
        with path.open("r", encoding="utf-8-sig", newline="") as fh:
            reader = csv.reader(fh, delimiter=delimiter, strict=True)
            width: int | None = None
            for row in reader:
                if width is None:
                    width = len(row)
                rows.append(list(row))
                _enforce_raw_dimensions(len(rows), max(width or 0, len(row)))
    except csv.Error as exc:
        raise MatrixFormatError(f"Malformed delimited file: {exc}") from exc
    return rows


def _enforce_xlsx_archive_budget(path: Path) -> None:
    """Bound decompressed XLSX payload before openpyxl allocates workbook objects."""
    try:
        with zipfile.ZipFile(path) as zf:
            total = 0
            for info in zf.infolist():
                if info.flag_bits & 0x1:
                    raise MatrixFormatError("Encrypted XLSX entries are unsupported.")
                if info.file_size > DEFAULT_MAX_XLSX_ENTRY_BYTES:
                    raise MatrixFormatError(
                        f"XLSX entry {info.filename!r} expands to {info.file_size:,} bytes, exceeding the "
                        f"per-entry cap of {DEFAULT_MAX_XLSX_ENTRY_BYTES:,} bytes."
                    )
                total += int(info.file_size)
                if total > DEFAULT_MAX_XLSX_UNCOMPRESSED_BYTES:
                    raise MatrixFormatError(
                        f"XLSX expands to more than {DEFAULT_MAX_XLSX_UNCOMPRESSED_BYTES:,} bytes, "
                        "exceeding the scalar-engine decompression cap."
                    )
    except zipfile.BadZipFile as exc:
        raise MatrixFormatError(f"Invalid XLSX archive: {exc}") from exc



def _excel_column_number(letters: str) -> int:
    value = 0
    for ch in letters.upper():
        if ch < "A" or ch > "Z":
            raise MatrixFormatError(f"Invalid XLSX cell reference column: {letters!r}.")
        value = value * 26 + (ord(ch) - ord("A") + 1)
    return value


def _xlsx_coordinate(ref: str) -> tuple[int, int]:
    match = _CELL_REF_RE.fullmatch(ref or "")
    if match is None:
        raise MatrixFormatError(f"Invalid XLSX cell reference: {ref!r}.")
    column = _excel_column_number(match.group(1))
    row = int(match.group(2))
    return row, column


_OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_WORKSHEET_REL_SUFFIX = "/worksheet"


def _xml_root_from_zip(zf: zipfile.ZipFile, name: str) -> ET.Element:
    try:
        data = zf.read(name)
    except KeyError as exc:
        raise MatrixFormatError(f"XLSX is missing required OOXML part {name!r}.") from exc
    try:
        return ET.fromstring(data)
    except ET.ParseError as exc:
        raise MatrixFormatError(f"Malformed OOXML part {name!r}: {exc}.") from exc


def _normalize_workbook_target(target: str) -> str:
    if not target:
        raise MatrixFormatError("XLSX workbook relationship has an empty worksheet target.")
    if target.startswith("/"):
        normalized = posixpath.normpath(target.lstrip("/"))
    else:
        normalized = posixpath.normpath(posixpath.join("xl", target))
    if normalized.startswith("../") or normalized == "..":
        raise MatrixFormatError(f"Unsupported XLSX worksheet relationship target: {target!r}.")
    return normalized


def _referenced_xlsx_worksheet_parts(zf: zipfile.ZipFile) -> list[str]:
    """Resolve the exact worksheet parts referenced by workbook relationships.

    The scalar importer deliberately supports only the conventional
    ``xl/worksheets/*.xml`` worksheet location. This makes the raw-XML preflight set
    identical to the worksheet set that openpyxl is permitted to load and closes
    relationship-target routing around the preflight boundary.
    """
    workbook = _xml_root_from_zip(zf, "xl/workbook.xml")
    rels = _xml_root_from_zip(zf, "xl/_rels/workbook.xml.rels")

    rel_by_id: dict[str, tuple[str, str, str | None]] = {}
    for rel in rels:
        if rel.tag.rsplit("}", 1)[-1] != "Relationship":
            continue
        rid = rel.attrib.get("Id")
        if not rid:
            continue
        rel_by_id[rid] = (
            rel.attrib.get("Type", ""),
            rel.attrib.get("Target", ""),
            rel.attrib.get("TargetMode"),
        )

    sheet_rids: list[str] = []
    rid_attr = f"{{{_OFFICE_REL_NS}}}id"
    for elem in workbook.iter():
        if elem.tag.rsplit("}", 1)[-1] == "sheet":
            rid = elem.attrib.get(rid_attr)
            if not rid:
                raise MatrixFormatError("XLSX workbook sheet is missing its relationship id.")
            sheet_rids.append(rid)

    if len(sheet_rids) != 1:
        raise MatrixFormatError(
            f"XLSX input must contain exactly one worksheet; found {len(sheet_rids)} workbook sheet entries."
        )

    parts: list[str] = []
    for rid in sheet_rids:
        if rid not in rel_by_id:
            raise MatrixFormatError(f"XLSX workbook sheet relationship {rid!r} is missing.")
        rel_type, target, target_mode = rel_by_id[rid]
        if target_mode and target_mode.lower() == "external":
            raise MatrixFormatError("External XLSX worksheet relationships are unsupported.")
        if not rel_type.endswith(_WORKSHEET_REL_SUFFIX):
            raise MatrixFormatError(
                f"XLSX sheet relationship {rid!r} is not a worksheet relationship ({rel_type!r})."
            )
        part = _normalize_workbook_target(target)
        # Strict supported-path contract: do not let openpyxl follow arbitrary OOXML
        # relationship targets that were not covered by worksheet preflight.
        if not (part.startswith("xl/worksheets/") and part.endswith(".xml")):
            raise MatrixFormatError(
                f"Unsupported XLSX worksheet relationship target {target!r}. "
                "Worksheets must use the standard xl/worksheets/*.xml location."
            )
        if part not in zf.namelist():
            raise MatrixFormatError(f"Referenced XLSX worksheet part {part!r} is missing from the archive.")
        parts.append(part)

    if len(set(parts)) != len(parts):
        raise MatrixFormatError("Multiple workbook sheet entries reference the same worksheet part.")
    return parts


def _preflight_xlsx_worksheet_xml(path: Path) -> None:
    """Reject ambiguous/unsafe worksheet structures before openpyxl object creation.

    The exact worksheet relationship target that openpyxl is allowed to load is first
    resolved from workbook.xml + workbook.xml.rels. Only the conventional
    xl/worksheets/*.xml location is supported. That referenced part is then streamed
    directly so formulas, merges, and sparse/far-away coordinates cannot be hidden by
    worksheet dimension metadata or relationship routing.
    """
    try:
        with zipfile.ZipFile(path) as zf:
            worksheet_names = _referenced_xlsx_worksheet_parts(zf)

            total_cell_records = 0
            max_row = 0
            max_col = 0
            for sheet_name in worksheet_names:
                with zf.open(sheet_name) as fh:
                    try:
                        for event, elem in ET.iterparse(fh, events=("start", "end")):
                            tag = elem.tag.rsplit("}", 1)[-1]
                            if event == "start" and tag == "mergeCell":
                                ref = elem.attrib.get("ref", "")
                                raise MatrixFormatError(
                                    f"Merged cells are unsupported in XLSX numeric matrices (found {ref!r}). "
                                    "Export an unmerged rectangular table before analysis."
                                )
                            if event == "start" and tag == "c":
                                ref = elem.attrib.get("r")
                                if not ref:
                                    raise MatrixFormatError("XLSX cell record is missing its coordinate reference.")
                                row, col = _xlsx_coordinate(ref)
                                total_cell_records += 1
                                max_row = max(max_row, row)
                                max_col = max(max_col, col)
                                _enforce_raw_dimensions(max_row, max_col)
                                if total_cell_records > DEFAULT_MAX_MATRIX_CELLS:
                                    raise MatrixFormatError(
                                        f"XLSX contains more than {DEFAULT_MAX_MATRIX_CELLS:,} cell records, "
                                        "exceeding the scalar-engine import cap."
                                    )
                            if event == "start" and tag == "row":
                                raw = elem.attrib.get("r")
                                if raw is not None:
                                    try:
                                        row = int(raw)
                                    except ValueError as exc:
                                        raise MatrixFormatError(f"Invalid XLSX row coordinate: {raw!r}.") from exc
                                    if row < 1:
                                        raise MatrixFormatError(f"Invalid XLSX row coordinate: {raw!r}.")
                                    max_row = max(max_row, row)
                                    _enforce_raw_dimensions(max_row, max_col)
                            if event == "start" and tag == "f":
                                raise MatrixFormatError(
                                    "XLSX formula cells are unsupported. Export literal values before analysis "
                                    "so identifiers and measurements are auditable."
                                )
                            if event == "end":
                                elem.clear()
                    except ET.ParseError as exc:
                        raise MatrixFormatError(f"Malformed worksheet XML in {sheet_name!r}: {exc}.") from exc
    except zipfile.BadZipFile as exc:
        raise MatrixFormatError(f"Invalid XLSX archive: {exc}") from exc

def _read_xlsx(path: Path) -> list[list[object]]:
    # Do not use read-only mode: it can trust incorrect producer dimensions. Before
    # normal-mode parsing, stream the raw worksheet XML to reject merges/formulas and
    # enforce actual coordinate/cell budgets. This prevents openpyxl merge processing
    # from discarding hidden cell records or allocating large merge rectangles first.
    _enforce_xlsx_archive_budget(path)
    _preflight_xlsx_worksheet_xml(path)
    wb = load_workbook(path, read_only=False, data_only=False)
    try:
        if len(wb.sheetnames) != 1:
            raise MatrixFormatError(
                f"XLSX input must contain exactly one worksheet; found {len(wb.sheetnames)}."
            )
        ws = wb[wb.sheetnames[0]]
        actual_rows = int(ws.max_row or 0)
        actual_columns = int(ws.max_column or 0)
        _enforce_raw_dimensions(actual_rows, actual_columns)
        rows: list[list[object]] = []
        for r_idx, row in enumerate(
            ws.iter_rows(min_row=1, max_row=actual_rows, min_col=1, max_col=actual_columns), start=1
        ):
            values: list[object] = []
            for c_idx, cell in enumerate(row, start=1):
                if getattr(cell, "data_type", None) == "f":
                    raise MatrixFormatError(
                        f"XLSX formula cell at row {r_idx}, column {c_idx} is unsupported. "
                        "Export literal values before analysis so identifiers and measurements are auditable."
                    )
                values.append(cell.value)
            rows.append(values)
        return rows
    finally:
        wb.close()


def load_matrix(
    path: str | Path,
    orientation: Orientation = "samples_rows",
    *,
    missing_tokens: Iterable[str] = DEFAULT_MISSING_TOKENS,
) -> pd.DataFrame:
    """Load a matrix and return samples × features while preserving identifiers exactly.

    Identifier cells never pass through pandas type inference. Measurement cells are
    parsed explicitly as real numeric values or configured missing tokens. Formula cells
    in XLSX are rejected in this reference implementation. XLSX content is preflighted from raw worksheet XML; merged cells and formulas are rejected before openpyxl parsing, and worksheet-dimension metadata is not trusted as a completeness bound.
    """
    path = Path(path)
    if not path.is_file():
        raise MatrixFormatError(f"Input file does not exist: {path}")
    size = path.stat().st_size
    if size > DEFAULT_MAX_FILE_BYTES:
        raise MatrixFormatError(
            f"Input file is {size:,} bytes, exceeding the scalar-engine import cap of {DEFAULT_MAX_FILE_BYTES:,} bytes."
        )
    suffix = path.suffix.lower()
    tokens = frozenset(str(x) for x in missing_tokens)
    if suffix == ".csv":
        rows = _read_delimited(path, ",")
    elif suffix in {".tsv", ".txt"}:
        rows = _read_delimited(path, "\t")
    elif suffix == ".xlsx":
        rows = _read_xlsx(path)
    elif suffix == ".xls":
        raise MatrixFormatError("Legacy .xls files are not supported. Convert to .xlsx, .csv, or .tsv.")
    else:
        raise MatrixFormatError(f"Unsupported file type {suffix!r}. Use .csv, .tsv, .txt, or .xlsx.")
    return _matrix_from_raw_rows(rows, orientation=orientation, missing_tokens=tokens)


def align_samples(a: pd.DataFrame, b: pd.DataFrame) -> tuple[pd.DataFrame, pd.DataFrame, dict[str, object]]:
    """Align matrices to one canonical exact-text sample order, preserving dataset A order."""
    for label, df in (("A", a), ("B", b)):
        if not df.index.is_unique:
            raise ValueError(f"Dataset {label} sample IDs must be unique before alignment.")
        if any(not isinstance(x, str) or x == "" for x in df.index):
            raise ValueError(f"Dataset {label} sample IDs must be nonempty strings.")

    b_ids = set(b.index)
    common = [sid for sid in a.index if sid in b_ids]
    only_a = [sid for sid in a.index if sid not in b_ids]
    a_ids = set(a.index)
    only_b = [sid for sid in b.index if sid not in a_ids]

    aligned_a = a.loc[common].copy()
    aligned_b = b.loc[common].copy()
    if not aligned_a.index.equals(aligned_b.index):
        raise AssertionError("Internal alignment invariant failed.")

    report = {
        "n_samples_a": int(len(a.index)),
        "n_samples_b": int(len(b.index)),
        "n_common": int(len(common)),
        "common_order": list(common),
        "canonical_order": list(common),
        "only_a": list(only_a),
        "only_b": list(only_b),
    }
    return aligned_a, aligned_b, report
