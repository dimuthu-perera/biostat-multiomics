from __future__ import annotations

import hashlib
import json
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook

from correlation_tool import (
    CorrelationConfig,
    DatasetDeclaration,
    MatrixFormatError,
    StudyDesignDeclaration,
    analyze_validated_run,
    create_acknowledgement,
    inspect_run,
    load_matrix,
)
from correlation_tool.analysis import _compute_cross_correlation

ROOT = Path(__file__).resolve().parents[2]
FIXTURES = Path(__file__).parent / "fixtures"
EXPECTED_HASHES = {
    "__init__.py": "af2da54b6fc28ad7b8c2e3a231aa0b6638f3031afe14de0632688ed22cb87f06",
    "analysis.py": "5f76c90e90662205aaf6ee62c739fc0f0419b65409f38ab111e89f87756faed8",
    "inspection.py": "1ab46306a85c3eff02371b63bf03b9214a1bae3d020c9f6fd081db8eee315ca0",
    "io.py": "ea24769efd2c801e0b3e3c584f11d1f19b79646137b176215d30edd0590d387d",
    "workflow.py": "f99b8eb107dff5544dce5a31316de9cc0307bcafee956e8f372956ade3df0df7",
}


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def _independent():
    return StudyDesignDeclaration("independent", "oracle fixture")


def _meta(name: str):
    return DatasetDeclaration("metabolomics", "quantitative", name)


def test_frozen_oracle_source_hashes_are_unchanged():
    for name, expected in EXPECTED_HASHES.items():
        assert _sha(ROOT / "correlation_tool" / name) == expected


@pytest.mark.parametrize("method", ["pearson", "spearman"])
def test_permanent_oracle_fixture_matches_accepted_v024_outputs(method):
    a = load_matrix(FIXTURES / "a.csv", "samples_rows")
    b = load_matrix(FIXTURES / "b.csv", "samples_rows")
    inspected = inspect_run(a, b, _meta("A"), _meta("B"), _independent(), CorrelationConfig(method=method))
    completed = analyze_validated_run(inspected, acknowledgement=create_acknowledgement(inspected))
    expected = json.loads((FIXTURES / f"expected_{method}.json").read_text())
    cols = ["feature_a", "feature_b", "method", "estimate", "p_value", "q_value", "n_pairwise", "status", "p_value_method", "permutations", "extreme_count"]
    actual = completed.results[cols].to_dict("records")
    assert len(actual) == len(expected)
    for got, want in zip(actual, expected, strict=True):
        for key in cols:
            if want[key] is None:
                assert pd.isna(got[key])
            elif isinstance(want[key], float):
                assert float(got[key]) == pytest.approx(want[key], abs=1e-14)
            else:
                assert got[key] == want[key]


@pytest.mark.parametrize(
    "x, expected_r, expected_p",
    [
        ([1e14, 1e14 + .01, 1e14 + .02, 1e14 + .03], 0.9486832980505138, 0.05131670194948623),
        ([1e-320, 2e-320, 3e-320, 4e-320], 1.0, 0.0),
    ],
)
def test_pearson_affine_stabilization_regression(x, expected_r, expected_p):
    idx = [f"S{i}" for i in range(4)]
    a = pd.DataFrame({"a": x}, index=idx)
    b = pd.DataFrame({"b": [1.0, 2.0, 3.0, 4.0]}, index=idx)
    result, _ = _compute_cross_correlation(a, b, CorrelationConfig(method="pearson"))
    assert result.loc[0, "estimate"] == pytest.approx(expected_r, abs=1e-15)
    assert result.loc[0, "p_value"] == pytest.approx(expected_p, abs=1e-15)


def _append_merge_range(path: Path, ref: str) -> None:
    tmp = path.with_suffix(".tmp.xlsx")
    with zipfile.ZipFile(path, "r") as zin, zipfile.ZipFile(tmp, "w") as zout:
        for info in zin.infolist():
            data = zin.read(info.filename)
            if info.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(data)
                ns = root.tag.split("}", 1)[0].strip("{")
                tag = f"{{{ns}}}"
                merges = root.find(f"{tag}mergeCells")
                if merges is None:
                    merges = ET.SubElement(root, f"{tag}mergeCells")
                ET.SubElement(merges, f"{tag}mergeCell", {"ref": ref})
                merges.set("count", str(len(list(merges))))
                data = ET.tostring(root, encoding="utf-8", xml_declaration=False)
            zout.writestr(info, data)
    tmp.replace(path)


def test_xlsx_merge_is_rejected_before_openpyxl(tmp_path, monkeypatch):
    import correlation_tool.io as cio

    p = tmp_path / "merge.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.append(["id", "x"])
    ws.append(["S0", 1])
    wb.save(p)
    _append_merge_range(p, "B2:B1001")
    called = False

    def sentinel(*args, **kwargs):
        nonlocal called
        called = True
        raise AssertionError("openpyxl must not run")

    monkeypatch.setattr(cio, "load_workbook", sentinel)
    with pytest.raises(MatrixFormatError, match="Merged cells are unsupported"):
        cio.load_matrix(p)
    assert called is False


def _relocate_worksheet(path: Path) -> None:
    tmp = path.with_suffix(".tmp.xlsx")
    rel_path = "xl/_rels/workbook.xml.rels"
    content_types_path = "[Content_Types].xml"
    source_sheet = "xl/worksheets/sheet1.xml"
    target = "/xl/data/matrix.xml"
    with zipfile.ZipFile(path, "r") as zin:
        sheet_data = zin.read(source_sheet)
        infos = zin.infolist()
        with zipfile.ZipFile(tmp, "w") as zout:
            for info in infos:
                data = zin.read(info.filename)
                if info.filename == rel_path:
                    root = ET.fromstring(data)
                    for rel in root:
                        if rel.attrib.get("Type", "").endswith("/worksheet"):
                            rel.set("Target", target)
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=False)
                elif info.filename == content_types_path:
                    root = ET.fromstring(data)
                    ns = root.tag.split("}", 1)[0].strip("{")
                    ET.SubElement(
                        root,
                        f"{{{ns}}}Override",
                        {
                            "PartName": target,
                            "ContentType": "application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml",
                        },
                    )
                    data = ET.tostring(root, encoding="utf-8", xml_declaration=False)
                zout.writestr(info, data)
            zout.writestr(target.lstrip("/"), sheet_data)
    tmp.replace(path)


@pytest.mark.parametrize("orientation", ["samples_rows", "samples_columns"])
def test_xlsx_nonstandard_relationship_target_is_rejected_before_openpyxl(tmp_path, monkeypatch, orientation):
    import correlation_tool.io as cio

    p = tmp_path / f"alt_{orientation}.xlsx"
    wb = Workbook()
    ws = wb.active
    if orientation == "samples_rows":
        ws.append(["id", "x"]); ws.append(["S0", 1]); ws.append(["S1", 2]); ws.append(["S2", 3])
    else:
        ws.append(["id", "S0", "S1", "S2"]); ws.append(["x", 1, 2, 3])
    wb.save(p)
    _relocate_worksheet(p)
    called = False

    def sentinel(*args, **kwargs):
        nonlocal called
        called = True
        raise AssertionError("openpyxl must not run")

    monkeypatch.setattr(cio, "load_workbook", sentinel)
    with pytest.raises(MatrixFormatError, match="Unsupported XLSX worksheet relationship target"):
        cio.load_matrix(p, orientation=orientation)
    assert called is False
